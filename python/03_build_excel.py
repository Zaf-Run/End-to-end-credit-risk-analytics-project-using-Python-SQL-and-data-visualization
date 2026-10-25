"""Step 3: build the Excel workbook with live formulas (openpyxl). Pivot tables are NOT created here: see the Pivot_Guide sheet."""
import pandas as pd, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
raw = pd.read_csv(f"{ROOT}/data/raw/bank.csv", sep=";"); n = len(raw); last = n + 1
F = "Arial"; hdr_fill = PatternFill("solid", fgColor="1F4E79"); inp_fill = PatternFill("solid", fgColor="FFFF00")
H = Font(name=F, bold=True, color="FFFFFF", size=10); B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True)
BLUE = Font(name=F, size=10, color="0000FF"); GREEN = Font(name=F, size=10, color="008000"); thin = Side(style="thin", color="BFBFBF")
def head(ws, row, vals, col=1):
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=col + j, value=v); c.font = H; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
wb = Workbook()
# ---------- Raw_Data
ws = wb.active; ws.title = "Raw_Data"; head(ws, 1, list(raw.columns))
for r in raw.itertuples(index=False): ws.append(list(r))
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = B
ws.freeze_panes = "A2"
for j in range(1, 18): ws.column_dimensions[L(j)].width = 13
# ---------- Lookups (inputs)
lk = wb.create_sheet("Lookups")
lk["A1"], lk["B1"] = "Indicator (illustrative, not validated)", "Points"; head(lk, 1, ["Indicator (illustrative, not validated)", "Points"])
lk["A2"], lk["B2"] = "Negative balance", 2; lk["A3"], lk["B3"] = "Personal loan", 1
head(lk, 1, ["Points", "Tier"], col=4)
for i, (p, t) in enumerate([(0, "Low"), (1, "Medium"), (2, "High"), (3, "High")]): lk.cell(row=2 + i, column=4, value=p); lk.cell(row=2 + i, column=5, value=t)
lk["A8"], lk["B8"] = "z value for 95% confidence interval", 1.96
for ref in ["B2", "B3", "B8"]: lk[ref].font = BLUE; lk[ref].fill = inp_fill
for row in lk.iter_rows(min_row=2):
    for c in row:
        if c.font != BLUE: c.font = B
lk["A10"] = "Blue on yellow cells are inputs. Change points or tiers and every sheet recalculates. Thresholds are illustrative rules chosen from risk logic, not statistically optimised cutoffs."
lk["A10"].font = Font(name=F, size=9, italic=True); lk.column_dimensions["A"].width = 42; lk.column_dimensions["D"].width = 10; lk.column_dimensions["E"].width = 12
# ---------- Cleaned_Data (formulas)
cd = wb.create_sheet("Cleaned_Data")
cols = ["customer_id", "age", "job", "marital", "education", "default_raw", "balance", "housing_loan", "personal_loan",
        "default_flag", "housing_loan_flag", "personal_loan_flag", "negative_balance_flag", "age_band", "balance_band", "risk_points", "risk_tier"]
head(cd, 1, cols)
src = {"age": "A", "job": "B", "marital": "C", "education": "D", "default_raw": "E", "balance": "F", "housing_loan": "G", "personal_loan": "H"}
for i in range(2, last + 1):
    cd.cell(row=i, column=1, value=i - 1)
    for j, name in enumerate(cols[1:9], start=2):
        s = src[name]; cd.cell(row=i, column=j, value=f'=SUBSTITUTE(Raw_Data!{s}{i},".","")' if name == "job" else f"=Raw_Data!{s}{i}")
    cd.cell(row=i, column=10, value=f'=IF(F{i}="yes",1,0)'); cd.cell(row=i, column=11, value=f'=IF(H{i}="yes",1,0)')
    cd.cell(row=i, column=12, value=f'=IF(I{i}="yes",1,0)'); cd.cell(row=i, column=13, value=f"=IF(G{i}<0,1,0)")
    cd.cell(row=i, column=14, value=f'=_xlfn.IFS(B{i}<30,"Under 30",B{i}<40,"30-39",B{i}<50,"40-49",B{i}<60,"50-59",TRUE,"60+")')
    cd.cell(row=i, column=15, value=f'=_xlfn.IFS(G{i}<0,"1 Negative",G{i}=0,"2 Zero",G{i}<=500,"3 1-500",G{i}<=2000,"4 501-2000",TRUE,"5 Above 2000")')
    cd.cell(row=i, column=16, value=f"=Lookups!$B$2*M{i}+Lookups!$B$3*L{i}")
    cd.cell(row=i, column=17, value=f"=INDEX(Lookups!$E$2:$E$5,MATCH(P{i},Lookups!$D$2:$D$5,0))")
for row in cd.iter_rows(min_row=2):
    for c in row: c.font = GREEN if c.column <= 9 else B
cd.freeze_panes = "A2"
for j in range(1, 18): cd.column_dimensions[L(j)].width = 15
# ---------- Data_Dictionary
dd = wb.create_sheet("Data_Dictionary")
head(dd, 1, ["Column", "Description", "Data type", "Business meaning", "Analytical relevance", "Missing values", "Potential quality issues"])
miss = raw.isna().sum()
unk = {c: int((raw[c] == "unknown").sum()) for c in ["job", "education", "contact", "poutcome"]}
rows = [
 ("age", "Customer age in years", "Integer", "Life stage proxy", "Descriptive only. Excluded from scoring (fairness)", "none", "Range 19 to 87. 38 IQR outliers, plausible values"),
 ("job", "Occupation category", "Text (12 levels)", "Employment type proxy", "Segment analysis. Employment status is not recorded directly", "none", f"{unk['job']} 'unknown'. Label 'admin.' cleaned to 'admin'"),
 ("marital", "Marital status", "Text (3 levels)", "Household context", "Descriptive only. Excluded from scoring (fairness)", "none", "None found"),
 ("education", "Education level", "Text (4 levels)", "Human capital proxy", "Descriptive only", "none", f"{unk['education']} 'unknown' kept as own category"),
 ("default", "Has credit in default (yes/no)", "Text yes/no", "Customer level default indicator. Contract definition not given", "TARGET for risk analysis", "none", "Only 76 yes. Definition and observation window unknown"),
 ("balance", "Average yearly balance in the account", "Integer", "Deposit account position, not loan exposure", "Strongest signal. May be a consequence of default, not a cause", "none", "Currency unit not stated. 366 negative. Max 71188, heavy right tail"),
 ("housing", "Has housing loan (yes/no)", "Text yes/no", "Mortgage product holder", "Exposure proxy (holder count only)", "none", "No amount, tenor or rate"),
 ("loan", "Has personal loan (yes/no)", "Text yes/no", "Unsecured credit holder", "Risk indicator and exposure proxy", "none", "No amount, tenor or rate"),
 ("contact", "Contact channel of last campaign call", "Text", "Marketing operations", "Excluded: not a credit variable", "none", f"{unk['contact']} 'unknown'"),
 ("day", "Day of month of last contact", "Integer", "Marketing operations", "Excluded", "none", "Not a default date"),
 ("month", "Month of last contact", "Text", "Marketing operations", "Excluded. Cannot support a default trend", "none", "Not a default date"),
 ("duration", "Last call length in seconds", "Integer", "Marketing operations", "Excluded: known only after the call (leakage for campaign target)", "none", "330 IQR outliers"),
 ("campaign", "Contacts made in this campaign", "Integer", "Marketing operations", "Excluded", "none", "Max 50"),
 ("pdays", "Days since previous campaign contact (-1 = never)", "Integer", "Marketing operations", "Excluded", "none", "-1 is a code, not a number of days"),
 ("previous", "Contacts before this campaign", "Integer", "Marketing operations", "Excluded", "none", "Max 25"),
 ("poutcome", "Outcome of previous campaign", "Text", "Marketing operations", "Excluded: not repayment behaviour", "none", f"{unk['poutcome']} 'unknown' (81.9%)"),
 ("y", "Subscribed to term deposit (yes/no)", "Text yes/no", "Marketing outcome", "Excluded: NOT a credit outcome", "none", "Do not confuse with default"),
 ("default_flag", "1 if default = yes", "Derived 0/1", "Target in numeric form", "Used in all rates", "0", "Formula in Cleaned_Data"),
 ("negative_balance_flag", "1 if balance below zero", "Derived 0/1", "Overdrawn indicator", "Risk indicator", "0", "May overlap with the default definition"),
 ("balance_band / age_band", "Banded versions", "Derived text", "Segments for monitoring", "Segmentation", "0", "Band edges are analyst choices"),
 ("risk_points / risk_tier", "Illustrative score and Low/Medium/High tier", "Derived", "Transparent rule based tiering", "Monitoring segmentation", "0", "Not validated on independent data"),
]
for r in rows: dd.append(list(r))
for row in dd.iter_rows(min_row=2):
    for c in row: c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
for j, w in enumerate([22, 38, 16, 34, 44, 12, 46], 1): dd.column_dimensions[L(j)].width = w
# ---------- Risk_Analysis
ra = wb.create_sheet("Risk_Analysis")
ra["A1"] = "Default rate by segment. Denominator: all customers in the segment. Defaults = default_flag = 1."; ra["A1"].font = BB
ra["A2"] = "Formula pattern: customers = COUNTIFS(segment column, segment). defaults = SUMIFS(default_flag, segment column, segment)."; ra["A2"].font = Font(name=F, size=9, italic=True)
head(ra, 4, ["Dimension", "Segment", "Customers", "Defaults", "Default rate", "95% CI low", "95% CI high", "Share of customers", "Share of defaults", "Lift vs portfolio", "Flag"])
rng = lambda col: f"Cleaned_Data!${col}$2:${col}${last}"
dims = [("balance_band", "O", ["1 Negative", "2 Zero", "3 1-500", "4 501-2000", "5 Above 2000"]), ("personal_loan", "I", ["yes", "no"]),
        ("housing_loan", "H", ["yes", "no"]), ("age_band", "N", ["Under 30", "30-39", "40-49", "50-59", "60+"]),
        ("job", "C", sorted(raw.job.str.replace(".", "", regex=False).unique())), ("marital", "D", ["married", "single", "divorced"]),
        ("education", "E", ["primary", "secondary", "tertiary", "unknown"]), ("risk_tier", "Q", ["Low", "Medium", "High"])]
r = 5; pos = {}
for name, col, segs in dims:
    pos[name] = (r, r + len(segs) - 1)
    for s in segs:
        ra.cell(row=r, column=1, value=name); ra.cell(row=r, column=2, value=s)
        ra.cell(row=r, column=3, value=f"=COUNTIFS({rng(col)},B{r})"); ra.cell(row=r, column=4, value=f"=SUMIFS({rng('J')},{rng(col)},B{r})")
        ra.cell(row=r, column=5, value=f"=IF(C{r}=0,0,D{r}/C{r})")
        z = "Lookups!$B$8"
        ra.cell(row=r, column=6, value=f"=IF(C{r}=0,0,(E{r}+{z}^2/(2*C{r})-{z}*SQRT(E{r}*(1-E{r})/C{r}+{z}^2/(4*C{r}^2)))/(1+{z}^2/C{r}))")
        ra.cell(row=r, column=7, value=f"=IF(C{r}=0,0,(E{r}+{z}^2/(2*C{r})+{z}*SQRT(E{r}*(1-E{r})/C{r}+{z}^2/(4*C{r}^2)))/(1+{z}^2/C{r}))")
        ra.cell(row=r, column=8, value=f"=C{r}/Summary!$B$4"); ra.cell(row=r, column=9, value=f"=D{r}/Summary!$B$5")
        ra.cell(row=r, column=10, value=f"=E{r}/Summary!$B$6")
        ra.cell(row=r, column=11, value=f'=IF(C{r}<100,"Small sample",IF(J{r}>=2,"High",IF(J{r}<0.5,"Low","Near portfolio")))')
        r += 1
    r += 1
ra_last = r
for row in ra.iter_rows(min_row=5, max_row=ra_last):
    for c in row:
        c.font = B
        if c.column in (5, 6, 7, 8, 9): c.number_format = "0.00%"
        if c.column == 10: c.number_format = "0.00"
        if c.column == 3 or c.column == 4: c.number_format = "#,##0"
ra.conditional_formatting.add(f"E5:E{ra_last}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="E67E73"))
ra.conditional_formatting.add(f"K5:K{ra_last}", CellIsRule(operator="equal", formula=['"High"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(name=F, bold=True, color="9C0006")))
ra.conditional_formatting.add(f"K5:K{ra_last}", CellIsRule(operator="equal", formula=['"Small sample"'], fill=PatternFill("solid", bgColor="FFF2CC")))
for j, w in enumerate([16, 16, 11, 10, 12, 11, 11, 14, 14, 13, 15], 1): ra.column_dimensions[L(j)].width = w
ra.freeze_panes = "A5"
# ---------- Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Portfolio summary (all figures are formulas)"; sm["A1"].font = Font(name=F, bold=True, size=12)
kp = [("Metric", "Value", "Definition"),
 ("Total customers", f"=COUNTA({rng('A')})", "Rows in Cleaned_Data"),
 ("Defaulted customers", f"=SUM({rng('J')})", "default = yes"),
 ("Overall default rate", "=B5/B4", "Defaults / customers"),
 ("95% CI low", "=(B6+Lookups!B8^2/(2*B4)-Lookups!B8*SQRT(B6*(1-B6)/B4+Lookups!B8^2/(4*B4^2)))/(1+Lookups!B8^2/B4)", "Wilson interval"),
 ("95% CI high", "=(B6+Lookups!B8^2/(2*B4)+Lookups!B8*SQRT(B6*(1-B6)/B4+Lookups!B8^2/(4*B4^2)))/(1+Lookups!B8^2/B4)", "Wilson interval"),
 ("Customers with negative balance", f"=SUM({rng('M')})", "balance < 0"),
 ("Default rate, negative balance", f'=SUMIFS({rng("J")},{rng("M")},1)/B9', "Defaults / negative balance customers"),
 ("Default rate, non negative balance", f'=SUMIFS({rng("J")},{rng("M")},0)/(B4-B9)', "Defaults / other customers"),
 ("Share of defaults from negative balance", f'=SUMIFS({rng("J")},{rng("M")},1)/B5', "Concentration"),
 ("Median balance, defaulted", f'=MEDIAN(IF({rng("J")}=1,{rng("G")}))', "Array formula"),
 ("Median balance, all customers", f"=MEDIAN({rng('G')})", ""),
 ("Customers with personal loan", f"=SUM({rng('L')})", "Holder count (no loan amounts available)"),
 ("Customers with housing loan", f"=SUM({rng('K')})", "Holder count (no loan amounts available)"),
 ("Average balance, all customers", f"=AVERAGE({rng('G')})", "Deposit balance, not loan exposure")]
for i, row in enumerate(kp, start=3):
    for j, v in enumerate(row, 1):
        sm.cell(row=i, column=j, value=v)
head(sm, 3, list(kp[0]))
# fix: header is row 3, first metric row 4
from openpyxl.worksheet.formula import ArrayFormula
sm["B13"] = ArrayFormula("B13", f'=MEDIAN(IF({rng("J")}=1,{rng("G")}))')
for row in sm.iter_rows(min_row=4, max_row=3 + len(kp) - 1):
    for c in row: c.font = B
for ref in ["B6", "B7", "B8", "B10", "B11", "B12"]: sm[ref].number_format = "0.00%"
for ref in ["B4", "B5", "B9", "B15", "B16"]: sm[ref].number_format = "#,##0"
sm["B13"].number_format = "#,##0"; sm["B14"].number_format = "#,##0"; sm["B17"].number_format = "#,##0"
sm["A20"] = "Reading the numbers"; sm["A20"].font = BB
sm["A21"] = '="Negative balance customers are "&TEXT(B9/B4,"0.0%")&" of the portfolio but hold "&TEXT(B12,"0.0%")&" of all defaults."'
sm["A22"] = '=IF(B10>2*B11,"Default rate for negative balance customers is more than twice the rate for others: priority monitoring segment.","No strong separation on this indicator.")'
sm["A23"] = "Association only. Balance may be a result of default or a shared third factor. No loan amounts exist, so exposure, LGD and ECL cannot be computed."
for a in ["A21", "A22", "A23"]: sm[a].font = B
sm.column_dimensions["A"].width = 42; sm.column_dimensions["B"].width = 16; sm.column_dimensions["C"].width = 48
# ---------- Dashboard
db = wb.create_sheet("Dashboard")
db["A1"] = "Credit Risk Monitoring Dashboard (Excel prototype)"; db["A1"].font = Font(name=F, bold=True, size=14, color="1F4E79")
tiles = [("Customers", "=Summary!B4", "#,##0"), ("Defaults", "=Summary!B5", "#,##0"), ("Default rate", "=Summary!B6", "0.00%"), ("Neg. balance default rate", "=Summary!B10", "0.0%")]
for i, (lab, f, fmt) in enumerate(tiles):
    c = 1 + i * 3
    db.cell(row=3, column=c, value=lab).font = Font(name=F, size=9, color="595959"); v = db.cell(row=4, column=c, value=f); v.font = Font(name=F, size=18, bold=True, color="1F4E79"); v.number_format = fmt
from openpyxl.chart.series import SeriesLabel
def chart(title, cat, vals, anchor, ytitle, pct=True, w=15, h=7.5):
    ch = BarChart(); ch.type = "col"; ch.title = title; ch.height = h; ch.width = w
    for ref, lab in vals:
        ch.add_data(ref, titles_from_data=False); ch.series[-1].tx = SeriesLabel(v=lab)
    ch.set_categories(cat); ch.y_axis.title = ytitle; ch.y_axis.numFmt = "0%"
    if len(vals) == 1: ch.legend = None
    else: ch.legend.position = "b"
    db.add_chart(ch, anchor)
b0, b1 = pos["balance_band"]; t0, t1 = pos["risk_tier"]; p0, p1 = pos["personal_loan"]
R_ = lambda c, a, b: Reference(ra, min_col=c, min_row=a, max_row=b)
chart("Default rate by balance band", R_(2, b0, b1), [(R_(5, b0, b1), "Default rate")], "A6", "Default rate")
chart("Default rate by risk tier", R_(2, t0, t1), [(R_(5, t0, t1), "Default rate")], "H6", "Default rate")
chart("Share of customers vs share of defaults (balance band)", R_(2, b0, b1), [(R_(8, b0, b1), "Share of customers"), (R_(9, b0, b1), "Share of defaults")], "A23", "Share")
chart("Default rate by personal loan", R_(2, p0, p1), [(R_(5, p0, p1), "Default rate")], "H23", "Default rate")
# header cells for chart series titles come from row b0-1 which is blank/header only for the first block: set explicit series title rows
db.column_dimensions["A"].width = 14
# ---------- Pivot guide
pg = wb.create_sheet("Pivot_Guide")
guide = ["Pivot tables must be created in Excel itself (this file was generated by code, which cannot create pivots).",
 "1. Go to Cleaned_Data. Click any cell. Insert > PivotTable > New Worksheet.",
 "2. Rows: balance_band. Values: Count of customer_id (rename 'Customers'), Sum of default_flag (rename 'Defaults').",
 "3. Add a calculated field: Default rate = default_flag / customer_id is awkward with counts, so instead set Values > Average of default_flag and format as percent. Average of a 0/1 flag equals the default rate.",
 "4. Repeat with risk_tier, job, age_band as Rows. Insert > Slicer on personal_loan and marital for filtering.",
 "5. Compare each pivot with the matching block in Risk_Analysis. The numbers must agree. That is your validation step.",
 "Lookup note: XLOOKUP works in Excel 365. This file uses INDEX/MATCH (risk_tier column) because it works in every version. XLOOKUP equivalent: =XLOOKUP(P2,Lookups!D2:D5,Lookups!E2:E5)."]
for i, t in enumerate(guide, 1): pg.cell(row=i, column=1, value=t).font = B
pg.column_dimensions["A"].width = 130
wb._sheets = [wb[s] for s in ["Raw_Data", "Cleaned_Data", "Data_Dictionary", "Lookups", "Risk_Analysis", "Summary", "Dashboard", "Pivot_Guide"]]
wb.save(f"{ROOT}/excel/credit_risk_workbook.xlsx"); print("saved", pos)
